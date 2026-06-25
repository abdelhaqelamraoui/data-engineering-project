"""Native PowerPoint diagrams (boxes + arrows), built from shapes - not
images - so they stay sharp at any zoom and match the deck's theme exactly.
"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from theme import BORDER, CHARCOAL, CYAN, GOLD, LIGHT_BG, MUTED_TEXT, TEAL, WHITE, _add_text


def _no_line(shape):
    shape.line.fill.background()
    shape.shadow.inherit = False


def box(slide, x, y, w, h, title, subtitle="", *, fill=WHITE, border=TEAL,
        title_color=CHARCOAL, title_size=12, sub_size=9):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.10
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(45000)
    tf.margin_right = Emu(45000)
    tf.margin_top = Emu(30000)
    tf.margin_bottom = Emu(30000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(title_size)
    run.font.bold = True
    run.font.color.rgb = title_color
    if subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(sub_size)
        run2.font.color.rgb = MUTED_TEXT
    return shp


def h_arrow(slide, x1, y, x2, color=MUTED_TEXT):
    w = x2 - x1
    h = Emu(110000)
    y0 = y - h // 2
    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1, y0, w, h)
    arr.fill.solid()
    arr.fill.fore_color.rgb = color
    _no_line(arr)
    try:
        arr.adjustments[0] = 0.55
        arr.adjustments[1] = 0.55
    except (IndexError, ValueError):
        pass
    return arr


def v_arrow(slide, x, y1, y2, color=MUTED_TEXT):
    h = y2 - y1
    w = Emu(110000)
    x0 = x - w // 2
    arr = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, x0, y1, w, h)
    arr.fill.solid()
    arr.fill.fore_color.rgb = color
    _no_line(arr)
    try:
        arr.adjustments[0] = 0.55
        arr.adjustments[1] = 0.55
    except (IndexError, ValueError):
        pass
    return arr


def label(slide, x, y, w, h, text, *, size=10, color=MUTED_TEXT, align=PP_ALIGN.CENTER, bold=False, italic=False):
    return _add_text(slide, x, y, w, h, text, size=size, color=color, align=align, bold=bold, italic=italic)


def architecture_diagram(slide, x, y, w, h):
    """Jetstream -> Producer -> Kafka -> {Spark -> HBase -> API -> Dashboard}
    with a parallel branch down to {Indexer -> Elasticsearch -> Kibana}.
    """
    n_top = 7
    gap = Emu(70000)
    box_w = Emu(int((w - gap * (n_top - 1)) / n_top))
    box_h = Emu(620000)
    row1_y = y
    row2_y = y + Emu(1750000)

    top_labels = [
        ("Jetstream", "wss firehose", WHITE, BORDER),
        ("Producer", "Python", WHITE, TEAL),
        ("Kafka", "bluesky-posts", WHITE, TEAL),
        ("Spark", "Structured Streaming", WHITE, TEAL),
        ("HBase", "REST :8080", WHITE, TEAL),
        ("API", "FastAPI", WHITE, TEAL),
        ("Dashboard", "Next.js", WHITE, GOLD),
    ]

    centers = []
    cx = x
    for i, (title, sub, fill, border) in enumerate(top_labels):
        box(slide, cx, row1_y, box_w, box_h, title, sub, fill=fill, border=border,
            title_size=11, sub_size=8)
        centers.append(cx + box_w // 2)
        if i < len(top_labels) - 1:
            h_arrow(slide, cx + box_w + Emu(8000), row1_y + box_h // 2, cx + box_w + gap - Emu(8000))
        cx += box_w + gap

    kafka_center_x = centers[2]

    bottom_labels = [
        ("Indexer", "Python", WHITE, CYAN),
        ("Elasticsearch", "single-node 8.15", WHITE, CYAN),
        ("Kibana", "web UI :5601", WHITE, CYAN),
    ]
    box_w2 = box_w
    bx = kafka_center_x - box_w2 // 2
    v_arrow(slide, kafka_center_x, row1_y + box_h + Emu(8000), row2_y - Emu(8000), color=CYAN)
    for i, (title, sub, fill, border) in enumerate(bottom_labels):
        box(slide, bx, row2_y, box_w2, box_h, title, sub, fill=fill, border=border,
            title_size=11, sub_size=8)
        if i < len(bottom_labels) - 1:
            h_arrow(slide, bx + box_w2 + Emu(8000), row2_y + box_h // 2, bx + box_w2 + gap - Emu(8000), color=CYAN)
        bx += box_w2 + gap

    label(slide, x, row1_y + box_h + Emu(70000), Emu(2400000), Emu(260000),
          "Trending pipeline", size=10, color=TEAL, align=PP_ALIGN.LEFT, bold=True)
    label(slide, kafka_center_x - Emu(1100000), row2_y + box_h + Emu(60000), Emu(2400000), Emu(260000),
          "Search & exploration (additive)", size=10, color=CYAN, align=PP_ALIGN.LEFT, bold=True)


def cleaning_pipeline_diagram(slide, x, y, w, h):
    """The shared preprocessing module's steps, as one left-to-right flow."""
    steps = [
        ("Raw post text", "from Kafka"),
        ("Lowercase +\nstrip URLs/mentions", "regex"),
        ("Extract hashtags", "#word -> term"),
        ("Strip punctuation\n& apostrophes", "don't -> dont"),
        ("Filter stopwords\n& short words", "len >= 3"),
        ("Terms", "hashtags + words"),
    ]
    n = len(steps)
    gap = Emu(60000)
    box_w = Emu(int((w - gap * (n - 1)) / n))
    box_h = Emu(900000)
    cy = y + (h - box_h) // 2

    cx = x
    for i, (title, sub) in enumerate(steps):
        is_last = i == len(steps) - 1
        b = box(slide, cx, cy, box_w, box_h, title, sub,
                fill=GOLD if is_last else WHITE,
                border=GOLD if is_last else TEAL,
                title_color=CHARCOAL, title_size=11, sub_size=8)
        if i < n - 1:
            h_arrow(slide, cx + box_w + Emu(6000), cy + box_h // 2, cx + box_w + gap - Emu(6000))
        cx += box_w + gap


def tech_grid(slide, x, y, w, h, items, *, cols=5):
    """items: list of (name, role) tuples, laid out in a grid of badges."""
    rows = -(-len(items) // cols)
    gap_x = Emu(120000)
    gap_y = Emu(140000)
    cell_w = Emu(int((w - gap_x * (cols - 1)) / cols))
    cell_h = Emu(int((h - gap_y * (rows - 1)) / rows))
    border_colors = [TEAL, CYAN, GOLD]
    for i, (name, role) in enumerate(items):
        r, c = divmod(i, cols)
        cx = x + c * (cell_w + gap_x)
        cy = y + r * (cell_h + gap_y)
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cell_w, cell_h)
        badge.adjustments[0] = 0.10
        badge.fill.solid()
        badge.fill.fore_color.rgb = WHITE
        badge.line.color.rgb = border_colors[i % len(border_colors)]
        badge.line.width = Pt(1.5)
        badge.shadow.inherit = False
        tf = badge.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Emu(40000)
        tf.margin_right = Emu(40000)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = name
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = CHARCOAL
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = role
        run2.font.size = Pt(9)
        run2.font.color.rgb = MUTED_TEXT


def schema_table(slide, x, y, w, h, rows):
    """rows: list of (table, row_key, columns, writer, reader) tuples."""
    n_rows = len(rows) + 1
    table_shape = slide.shapes.add_table(n_rows, 5, x, y, w, h)
    table = table_shape.table
    headers = ["Table", "Row key", "Columns", "Written by", "Read by"]
    widths = [0.16, 0.24, 0.28, 0.16, 0.16]
    for i, frac in enumerate(widths):
        table.columns[i].width = Emu(int(w * frac))

    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CHARCOAL
        cell.margin_top = Emu(40000)
        cell.margin_bottom = Emu(40000)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = htext
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = WHITE

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT_BG
            cell.margin_top = Emu(30000)
            cell.margin_bottom = Emu(30000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = value
            run.font.size = Pt(10.5)
            run.font.color.rgb = CHARCOAL
            run.font.bold = c == 0
    return table_shape


def window_scoring_diagram(slide, x, y, w, h):
    """Sliding window timeline (top) + the ratio-to-baseline formula (bottom)."""
    timeline_h = Emu(70000)
    timeline_y = y + Emu(150000)
    timeline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, timeline_y, w, timeline_h)
    timeline.fill.solid()
    timeline.fill.fore_color.rgb = BORDER
    _no_line(timeline)

    win_w = Emu(int(w * 0.32))
    win_h = Emu(420000)
    step = Emu(int((w - win_w) / 3))
    colors = [TEAL, CYAN, GOLD]
    for i in range(3):
        wx = x + step * i
        wy = timeline_y - win_h // 2 + timeline_h // 2 + Emu(i * 90000)
        win = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, wx, wy, win_w, win_h)
        win.adjustments[0] = 0.15
        win.fill.solid()
        win.fill.fore_color.rgb = colors[i]
        win.line.color.rgb = WHITE
        win.line.width = Pt(1.5)
        win.shadow.inherit = False
        tf = win.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f"window {i + 1}"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = WHITE if colors[i] != GOLD else CHARCOAL

    label(slide, x, timeline_y + win_h + Emu(260000), w, Emu(300000),
          "5-minute window  ·  1-minute slide  ·  2-minute watermark",
          size=12, color=MUTED_TEXT, align=PP_ALIGN.CENTER, italic=True)

    formula_y = timeline_y + win_h + Emu(700000)
    formula_h = Emu(900000)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Emu(int(w * 0.18)), formula_y,
                                    Emu(int(w * 0.64)), formula_h)
    card.adjustments[0] = 0.12
    card.fill.solid()
    card.fill.fore_color.rgb = CHARCOAL
    _no_line(card)
    tf = card.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "score = count / (baseline_avg + smoothing)"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = GOLD
    r.font.name = "Consolas"

    label(slide, x, formula_y + formula_h + Emu(120000), w, Emu(300000),
          "Ratio to the term's own trailing baseline (EMA) - acceleration, not raw frequency",
          size=12, color=MUTED_TEXT, align=PP_ALIGN.CENTER)


def before_after_diagram(slide, x, y, w, h, *, raw_text, cleaned_text, hashtags, words):
    """One concrete post run through the real shared preprocess() function,
    shown as two side-by-side cards plus the extracted term chips.
    """
    col_gap = Emu(260000)
    col_w = (w - col_gap) // 2
    card_h = h

    before = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, col_w, card_h)
    before.adjustments[0] = 0.04
    before.fill.solid()
    before.fill.fore_color.rgb = WHITE
    before.line.color.rgb = BORDER
    before.line.width = Pt(1.25)
    before.shadow.inherit = False

    after_x = x + col_w + col_gap
    after = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, after_x, y, col_w, card_h)
    after.adjustments[0] = 0.04
    after.fill.solid()
    after.fill.fore_color.rgb = WHITE
    after.line.color.rgb = TEAL
    after.line.width = Pt(1.5)
    after.shadow.inherit = False

    pad = Emu(180000)
    label(slide, x + pad, y + Emu(130000), col_w - 2 * pad, Emu(260000),
          "BEFORE · RAW TEXT", size=10, bold=True, color=MUTED_TEXT, align=PP_ALIGN.LEFT)
    label(slide, x + pad, y + Emu(450000), col_w - 2 * pad, Emu(1100000),
          raw_text, size=13, color=CHARCOAL, align=PP_ALIGN.LEFT)

    label(slide, after_x + pad, y + Emu(130000), col_w - 2 * pad, Emu(260000),
          "AFTER · CLEANED TEXT", size=10, bold=True, color=TEAL, align=PP_ALIGN.LEFT)
    label(slide, after_x + pad, y + Emu(450000), col_w - 2 * pad, Emu(800000),
          cleaned_text, size=13, color=CHARCOAL, align=PP_ALIGN.LEFT)

    chip_y = y + Emu(1450000)
    chip_h = Emu(330000)
    cx = after_x + pad
    max_x = after_x + col_w - pad
    for tag in hashtags:
        chip_w = Emu(int(120000 + 65000 * len(tag)))
        if cx + chip_w > max_x:
            cx = after_x + pad
            chip_y += chip_h + Emu(80000)
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, chip_y, chip_w, chip_h)
        chip.adjustments[0] = 0.5
        chip.fill.solid()
        chip.fill.fore_color.rgb = TEAL
        _no_line(chip)
        tf = chip.text_frame
        tf.margin_left = Emu(10000)
        tf.margin_right = Emu(10000)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = tag
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = WHITE
        cx += chip_w + Emu(70000)
    for word in words:
        chip_w = Emu(int(120000 + 60000 * len(word)))
        if cx + chip_w > max_x:
            cx = after_x + pad
            chip_y += chip_h + Emu(80000)
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, chip_y, chip_w, chip_h)
        chip.adjustments[0] = 0.5
        chip.fill.solid()
        chip.fill.fore_color.rgb = WHITE
        chip.line.color.rgb = BORDER
        chip.line.width = Pt(0.75)
        chip.shadow.inherit = False
        tf = chip.text_frame
        tf.margin_left = Emu(10000)
        tf.margin_right = Emu(10000)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = word
        r.font.size = Pt(10)
        r.font.color.rgb = CHARCOAL
        cx += chip_w + Emu(70000)
